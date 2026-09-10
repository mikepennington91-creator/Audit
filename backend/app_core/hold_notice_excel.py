"""Build the controlled hold-notice register workbook."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


HOLD_NOTICE_COLUMNS = (
    ("Reference No", "reference", "text"),
    ("Date", "event_date", "date"),
    ("Week No.", "week_number", "number"),
    ("Year", "year", "number"),
    ("Name", "created_by_name", "text"),
    ("Product Code", "product_code", "text"),
    ("Product", "ingredient_name", "text"),
    ("Line Number", "line_area", "text"),
    ("Batch & Best before", "batch_and_best_before", "text"),
    ("Quantity (Cases / KG)", "quantity", "text"),
    ("Pallet Numbers (Separate multiple ID's with / )", "pallet_numbers", "text"),
    ("Reason for hold", "reason", "text"),
    ("Action taken / Required", "action_required", "text"),
    ("Root Cause", "root_cause", "text"),
    ("Root cause category", "root_cause_category", "text"),
    ("Conclusion", "conclusion", "text"),
    ("Quantity Rejected", "quantity_rejected", "text"),
    ("Quantity Released", "quantity_released", "text"),
    ("Date Closed", "date_closed", "date"),
    ("Status2", "status", "text"),
    ("Disposal (Y/N)", "has_disposal", "text"),
    ("Disposal Route", "disposal_route", "text"),
    ("Authorised by", "authorised_by", "text"),
    ("Disposal Reason", "disposal_reason", "text"),
)


def _safe_text(value: Any) -> str:
    if isinstance(value, (list, tuple, set)):
        value = " / ".join(str(item).strip() for item in value if str(item).strip())
    text = "" if value is None else str(value).strip()
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _date_value(value: Any) -> date | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for candidate, pattern in ((text[:10], "%Y-%m-%d"), (text[:10], "%d/%m/%Y")):
        try:
            return datetime.strptime(candidate, pattern).date()
        except ValueError:
            continue
    return None


def _batch_summary(hold: dict[str, Any]) -> str:
    parts = []
    if hold.get("our_batch"):
        parts.append(str(hold["our_batch"]).strip())
    if hold.get("vendor_batch"):
        parts.append(f"Vendor: {str(hold['vendor_batch']).strip()}")
    best_before = hold.get("best_before") or hold.get("best_before_date")
    if best_before:
        parsed = _date_value(best_before)
        parts.append(f"BBE: {parsed.strftime('%d/%m/%Y') if parsed else str(best_before).strip()}")
    return " / ".join(parts)


def _pallet_summary(hold: dict[str, Any]) -> str:
    value = hold.get("pallet_numbers") or hold.get("pallet_ids") or hold.get("pallet_number") or ""
    return _safe_text(value)


def hold_notice_export_row(hold: dict[str, Any], disposal: dict[str, Any] | None = None) -> dict[str, Any]:
    held_on = _date_value(hold.get("event_date"))
    outcome_present = any(hold.get(field) not in (None, "") for field in (
        "quantity_released", "quantity_discarded", "root_cause", "corrective_action"
    ))
    closed_on = _date_value(hold.get("date_closed") or hold.get("outcome_updated_at"))
    if not closed_on and disposal:
        closed_on = _date_value(disposal.get("event_date") or disposal.get("created_at"))
    conclusion = hold.get("conclusion") or (disposal or {}).get("action_required") or hold.get("corrective_action") or ""
    status = hold.get("status") or ("Closed" if outcome_present or disposal else "Open")
    return {
        "reference": hold.get("reference"),
        "event_date": held_on,
        "week_number": held_on.isocalendar().week if held_on else None,
        "year": held_on.year if held_on else None,
        "created_by_name": hold.get("created_by_name"),
        "product_code": hold.get("product_code") or hold.get("rm_number"),
        "ingredient_name": hold.get("product") or hold.get("ingredient_name"),
        "line_area": hold.get("line_number") or hold.get("line_area"),
        "batch_and_best_before": _batch_summary(hold),
        "quantity": hold.get("quantity"),
        "pallet_numbers": _pallet_summary(hold),
        "reason": hold.get("reason"),
        "action_required": hold.get("action_required"),
        "root_cause": hold.get("root_cause"),
        "root_cause_category": hold.get("root_cause_category"),
        "conclusion": conclusion,
        "quantity_rejected": hold.get("quantity_rejected") or hold.get("quantity_discarded"),
        "quantity_released": hold.get("quantity_released"),
        "date_closed": closed_on,
        "status": status,
        "has_disposal": "Y" if disposal else "N",
        "disposal_route": (disposal or {}).get("disposal_route_label") or (disposal or {}).get("disposal_route"),
        "authorised_by": (disposal or {}).get("authorised_by") or (disposal or {}).get("created_by_name") or hold.get("authorised_by"),
        "disposal_reason": (disposal or {}).get("reason"),
    }


def build_hold_notice_workbook(
    holds: Iterable[dict[str, Any]],
    disposals_by_hold_id: dict[str, dict[str, Any]] | None = None,
) -> bytes:
    """Return a filterable Excel register populated from hold and disposal records."""
    disposal_lookup = disposals_by_hold_id or {}
    rows = [hold_notice_export_row(hold, disposal_lookup.get(str(hold.get("id")))) for hold in holds]

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Hold Notices"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:X{max(2, len(rows) + 1)}"

    header_fill = PatternFill("solid", fgColor="991B1B")
    for column_number, (label, _field, _kind) in enumerate(HOLD_NOTICE_COLUMNS, start=1):
        cell = worksheet.cell(row=1, column=column_number, value=label)
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 45

    for row_number, record in enumerate(rows, start=2):
        for column_number, (_label, field, kind) in enumerate(HOLD_NOTICE_COLUMNS, start=1):
            value = record.get(field)
            if kind == "date":
                value = _date_value(value)
            elif kind == "text":
                value = _safe_text(value)
            cell = worksheet.cell(row=row_number, column=column_number, value=value)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=field in {
                "batch_and_best_before", "pallet_numbers", "reason", "action_required",
                "root_cause", "conclusion", "disposal_reason",
            })
            if kind == "date" and value:
                cell.number_format = "dd/mm/yyyy"
            elif kind == "number":
                cell.number_format = "0"

    widths = {
        "A": 16, "B": 13, "C": 10, "D": 10, "E": 21, "F": 17, "G": 28, "H": 18,
        "I": 28, "J": 22, "K": 28, "L": 42, "M": 42, "N": 38, "O": 24, "P": 42,
        "Q": 20, "R": 20, "S": 14, "T": 13, "U": 15, "V": 22, "W": 21, "X": 38,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    last_row = max(1, len(rows) + 1)
    table = Table(displayName="HoldNoticeRegister", ref=f"A1:X{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    worksheet.add_table(table)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
