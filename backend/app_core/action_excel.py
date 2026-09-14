"""Build the corrective-action register workbook."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ACTION_COLUMNS = (
    ("Reference", "reference", "text"),
    ("Action", "audit_name", "text"),
    ("Source", "source", "text"),
    ("Audit Question", "question_text", "text"),
    ("Non-Conformance", "non_conformance", "text"),
    ("Action Required", "action_required", "text"),
    ("Action Owner", "assigned_to", "text"),
    ("Raised By", "created_by_name", "text"),
    ("Approver", "reviewer_name", "text"),
    ("Date Raised", "created_at", "datetime"),
    ("Due Date", "due_date", "date"),
    ("Status", "status", "text"),
    ("Action Taken", "action_taken", "text"),
    ("Completed By", "completed_by_name", "text"),
    ("Date Completed", "completed_at", "datetime"),
    ("Review Comment", "review_comment", "text"),
    ("Reviewed By", "reviewed_by_name", "text"),
    ("Date Reviewed", "reviewed_at", "datetime"),
    ("Effectiveness Evidence", "effectiveness_evidence", "text"),
    ("Effectiveness Verified By", "effectiveness_verified_by_name", "text"),
    ("Effectiveness Verified Date", "effectiveness_verified_at", "datetime"),
    ("Archived", "archived", "text"),
)


def _safe_text(value: Any) -> str:
    text = "" if value is None else str(value).strip()
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _date_value(value: Any) -> date | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).date()
    except ValueError:
        return None


def _datetime_value(value: Any) -> datetime | None:
    if not value:
        return None
    if isinstance(value, datetime):
        parsed = value
    else:
        try:
            parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except ValueError:
            return None
    return parsed.replace(tzinfo=None)


def action_export_row(action: dict[str, Any]) -> dict[str, Any]:
    return {
        **action,
        "source": "Manual" if not action.get("run_id") else "Audit",
        "assigned_to": action.get("assigned_user_name") or action.get("assigned_department") or "Unassigned",
        "reviewer_name": action.get("reviewer_user_name") or action.get("created_by_name"),
        "status": str(action.get("status") or "open").replace("_", " ").title(),
        "archived": "Yes" if action.get("archived") else "No",
    }


def build_action_workbook(actions: Iterable[dict[str, Any]]) -> bytes:
    """Return a formatted, filterable Excel register for corrective actions."""
    rows = [action_export_row(action) for action in actions]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Corrective Actions"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="0F766E")
    for column_number, (label, _field, _kind) in enumerate(ACTION_COLUMNS, start=1):
        cell = worksheet.cell(row=1, column=column_number, value=label)
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 36

    wrapped_fields = {
        "question_text", "non_conformance", "action_required", "action_taken",
        "review_comment", "effectiveness_evidence",
    }
    for row_number, record in enumerate(rows, start=2):
        for column_number, (_label, field, kind) in enumerate(ACTION_COLUMNS, start=1):
            value = record.get(field)
            if kind == "date":
                value = _date_value(value)
            elif kind == "datetime":
                value = _datetime_value(value)
            else:
                value = _safe_text(value)
            cell = worksheet.cell(row=row_number, column=column_number, value=value)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=field in wrapped_fields)
            if kind == "date" and value:
                cell.number_format = "dd/mm/yyyy"
            elif kind == "datetime" and value:
                cell.number_format = "dd/mm/yyyy hh:mm"

    widths = {
        "A": 13, "B": 28, "C": 12, "D": 34, "E": 42, "F": 42, "G": 22,
        "H": 22, "I": 22, "J": 18, "K": 14, "L": 20, "M": 42, "N": 22,
        "O": 18, "P": 34, "Q": 22, "R": 18, "S": 42, "T": 25, "U": 21,
        "V": 12,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    last_row = max(1, len(rows) + 1)
    table = Table(displayName="CorrectiveActionRegister", ref=f"A1:V{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    worksheet.add_table(table)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
