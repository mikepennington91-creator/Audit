import asyncio
import io
import os
import sys
from datetime import date, time
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

os.environ.setdefault("DATABASE_URL", "postgresql://test:test@localhost/test")
os.environ.setdefault("JWT_SECRET_KEY", "test-secret")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import server
from app_core import document_imports, documents


@pytest.fixture
def template():
    return {
        "id": "template-1",
        "title": "Chocolate Line Check",
        "document_reference": "ZRO-PROD-01",
        "version": 3,
        "authorised_by": "QA Manager",
        "company_id": "company-1",
        "fields": [
            {"id": "batch", "label": "Batch Code", "field_type": "text", "section": "header", "required": True, "order": 0},
            {"id": "made", "label": "Production Date", "field_type": "date", "section": "header", "required": True, "order": 1},
            {"id": "decision", "label": "Decision", "field_type": "dropdown", "section": "header", "dropdown_options": ["Release", "Hold"], "order": 2},
            {"id": "check_time", "label": "Check Time", "field_type": "time", "section": "table", "required": True, "order": 3},
            {"id": "weight", "label": "Weight", "field_type": "number", "section": "table", "order": 4},
        ],
    }


def test_extraction_schema_is_strict_and_template_specific(template):
    schema = document_imports.extraction_schema(template["fields"])
    assert schema["additionalProperties"] is False
    assert set(schema["properties"]["header_values"]["required"]) == {"batch", "made", "decision"}
    row = schema["properties"]["table_rows"]["items"]
    assert set(row["properties"]["values"]["required"]) == {"check_time", "weight"}


@pytest.mark.parametrize("content,expected", [
    (b"%PDF-1.7 sample", "application/pdf"),
    (b"\xff\xd8\xff\xe0 sample", "image/jpeg"),
    (b"\x89PNG\r\n\x1a\n sample", "image/png"),
    (b"RIFF\x00\x00\x00\x00WEBP sample", "image/webp"),
    (b"not an image", None),
])
def test_scan_type_is_detected_from_file_content(content, expected):
    assert document_imports._detect_content_type(content) == expected


def test_normalise_extraction_flags_uncertain_and_invalid_values(template):
    values = document_imports.normalise_extraction(template, {
        "document_quality": "usable",
        "warnings": ["Bottom edge is slightly cropped"],
        "header_values": {"batch": " LOT 42 ", "made": "08/09/2026", "decision": "unknown"},
        "header_confidence": {"batch": 0.99, "made": 0.9, "decision": 0.4},
        "table_rows": [{
            "values": {"check_time": "09:15", "weight": 31.2},
            "confidence": {"check_time": 0.98, "weight": 0.8},
        }],
    })
    assert {value["field_id"]: value["value"] for value in values["field_values"]} == {
        "batch": "LOT 42", "made": "2026-09-08", "decision": "",
    }
    assert values["table_rows"] == [{"check_time": "09:15", "weight": "31.2"}]
    assert {flag["field_id"] for flag in values["review_flags"]} == {"decision", "weight"}
    assert any("allowed option" in warning for warning in values["warnings"])


def test_generated_workbook_round_trips_records_and_orders_table_rows(template):
    content = document_imports.build_import_workbook(template)
    workbook = load_workbook(io.BytesIO(content))
    assert workbook["_Infinit"].sheet_state == "hidden"
    assert workbook["Documents"].freeze_panes == "A3"
    documents_sheet = workbook["Documents"]
    documents_sheet.append(["DOC-001", "LOT 42", date(2026, 9, 8), "Release"])
    rows = workbook["Table Rows"]
    rows.append(["DOC-001", 2, time(10, 30), 31.4])
    rows.append(["DOC-001", 1, time(9, 15), 31.2])
    stream = io.BytesIO()
    workbook.save(stream)

    parsed = document_imports.parse_import_workbook(stream.getvalue(), template)
    assert parsed[0]["import_id"] == "DOC-001"
    assert {value["field_id"]: value["value"] for value in parsed[0]["field_values"]}["made"] == "2026-09-08"
    assert parsed[0]["table_rows"] == [
        {"check_time": "09:15", "weight": 31.2},
        {"check_time": "10:30", "weight": 31.4},
    ]


def test_workbook_rejects_wrong_template_version(template):
    workbook = load_workbook(io.BytesIO(document_imports.build_import_workbook(template)))
    workbook["_Infinit"]["B2"] = 2
    stream = io.BytesIO()
    workbook.save(stream)
    with pytest.raises(ValueError, match="changed"):
        document_imports.parse_import_workbook(stream.getvalue(), template)


def test_workbook_rejects_invalid_mapped_dropdown(template):
    workbook = load_workbook(io.BytesIO(document_imports.build_import_workbook(template)))
    workbook["Documents"].append(["DOC-001", "LOT 42", date(2026, 9, 8), "Destroy"])
    stream = io.BytesIO()
    workbook.save(stream)
    with pytest.raises(ValueError, match="Decision has an invalid dropdown"):
        document_imports.parse_import_workbook(stream.getvalue(), template)


def test_imported_document_is_always_created_as_unpublished_draft(template, monkeypatch):
    inserted = []

    class Collection:
        async def insert_one(self, row):
            inserted.append(dict(row))

    monkeypatch.setattr(server, "db", SimpleNamespace(traceability_documents=Collection()))
    user = {"id": "user-1", "name": "Reviewer", "company_id": "company-1"}
    result = asyncio.run(document_imports._create_imported_draft(
        template, user, {"field_values": [], "table_rows": []}, {"type": "excel", "import_id": "DOC-1"},
    ))
    assert result["completed"] is False
    assert result["completed_at"] is None
    assert result["imported_draft"] is True
    assert inserted[0]["import_source"]["import_id"] == "DOC-1"


def test_publishing_imported_draft_records_human_reviewer(monkeypatch):
    source_updates = []

    class Collection:
        async def update_one(self, query, update):
            source_updates.append((query, update["$set"]))

    imported = {"id": "doc-1", "completed": False, "imported_draft": True, "import_source": {"import_item_id": "item-1"}}
    monkeypatch.setattr(documents, "_get_accessible_document", lambda *args, **kwargs: asyncio.sleep(0, result=imported))
    monkeypatch.setattr(server, "update_traceability_document", lambda *args, **kwargs: asyncio.sleep(0, result={**imported, "completed": True}))
    monkeypatch.setattr(server, "db", SimpleNamespace(traceability_documents=Collection(), document_import_items=Collection()))
    actor = {"id": "reviewer-1", "name": "QA Reviewer"}
    data = server.TraceabilityDocumentSubmit(field_values=[], table_rows=[], completed=True)
    result = asyncio.run(documents.update_traceability_document("doc-1", data, actor))
    assert result["import_status"] == "published"
    assert result["import_reviewed_by"] == "reviewer-1"
    assert any(query == {"id": "item-1"} and values["status"] == "published" for query, values in source_updates)


def test_admin_close_out_cannot_bypass_import_review(monkeypatch):
    imported = {"id": "doc-1", "completed": False, "imported_draft": True}
    monkeypatch.setattr(documents, "_get_accessible_document", lambda *args, **kwargs: asyncio.sleep(0, result=imported))
    actor = {"id": "admin-1", "name": "Admin", "role": "admin", "company_id": "company-1"}
    with pytest.raises(Exception) as error:
        asyncio.run(documents.close_out_traceability_document("doc-1", actor))
    assert getattr(error.value, "status_code", None) == 409
