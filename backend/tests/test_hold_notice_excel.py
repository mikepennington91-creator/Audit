"""Coverage for the hold-notice register Excel export."""

from io import BytesIO

from openpyxl import load_workbook

from app_core.hold_notice_excel import HOLD_NOTICE_COLUMNS, build_hold_notice_workbook, hold_notice_export_row


def hold_record(**changes):
    record = {
        "id": "hold-1",
        "reference": "26-001",
        "event_date": "2026-03-18",
        "created_by_name": "Mike Pennington",
        "rm_number": "RM100",
        "ingredient_name": "Rework CDM Heroes Crunchie",
        "line_area": "Warehouse",
        "our_batch": "OBO0460262",
        "best_before_date": "2026-03-12",
        "quantity": "180 kg",
        "pallet_numbers": ["P1", "P2"],
        "reason": "Boxes are double labelled with differing weights.",
        "action_required": "Contact supplier",
        "root_cause": "Supplier did not label pallets correctly",
        "root_cause_category": "Supplier Issue",
        "quantity_discarded": "180 kg",
        "quantity_released": "0 kg",
        "outcome_updated_at": "2026-03-26T10:00:00+00:00",
    }
    return {**record, **changes}


def disposal_record(**changes):
    record = {
        "source_hold_id": "hold-1",
        "event_date": "2026-03-26",
        "disposal_route_label": "SugaRich",
        "created_by_name": "Ciaran Farren",
        "reason": "Incorrect supplier pallet labels",
        "action_required": "Send product to SugaRich for disposal",
    }
    return {**record, **changes}


def test_export_row_populates_hold_outcome_and_linked_disposal_fields():
    row = hold_notice_export_row(hold_record(), disposal_record())
    assert row["week_number"] == 12
    assert row["year"] == 2026
    assert row["product_code"] == "RM100"
    assert row["batch_and_best_before"] == "OBO0460262 / BBE: 12/03/2026"
    assert row["pallet_numbers"] == "P1 / P2"
    assert row["quantity_rejected"] == "180 kg"
    assert row["date_closed"].isoformat() == "2026-03-26"
    assert row["status"] == "Closed"
    assert row["has_disposal"] == "Y"
    assert row["disposal_route"] == "SugaRich"
    assert row["authorised_by"] == "Ciaran Farren"
    assert row["conclusion"] == "Send product to SugaRich for disposal"


def test_hold_without_outcome_or_disposal_is_open():
    row = hold_notice_export_row(hold_record(
        quantity_discarded="", quantity_released="", root_cause="",
        root_cause_category="", outcome_updated_at="",
    ))
    assert row["status"] == "Open"
    assert row["has_disposal"] == "N"
    assert row["date_closed"] is None


def test_workbook_has_requested_headers_filters_real_dates_and_safe_text():
    content = build_hold_notice_workbook(
        [hold_record(reason="=HYPERLINK(\"bad\")")],
        {"hold-1": disposal_record()},
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["Hold Notices"]
    assert [cell.value for cell in sheet[1]] == [label for label, _field, _kind in HOLD_NOTICE_COLUMNS]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:X2"
    assert sheet["B2"].value.isoformat() == "2026-03-18T00:00:00"
    assert sheet["B2"].number_format == "dd/mm/yyyy"
    assert sheet["L2"].value.startswith("'=")
    assert sheet.tables["HoldNoticeRegister"].ref == "A1:X2"
    workbook.close()
