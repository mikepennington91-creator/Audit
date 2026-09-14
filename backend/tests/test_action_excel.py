import os
import sys
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


os.environ.setdefault("DATABASE_URL", "postgresql://test:test@localhost:5432/test")
os.environ.setdefault("JWT_SECRET_KEY", "test-secret")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app_core.action_excel import build_action_workbook  # noqa: E402


def test_action_workbook_contains_filterable_register_fields():
    workbook_bytes = build_action_workbook([{
        "audit_name": "Glass audit",
        "run_id": "run-1",
        "question_text": "Is the guard intact?",
        "non_conformance": "Guard damaged",
        "action_required": "Replace guard",
        "assigned_user_name": "Chris",
        "created_by_name": "Mike",
        "reviewer_user_name": "Ciaran",
        "created_at": "2026-09-14T09:30:00+01:00",
        "due_date": "2026-09-21",
        "status": "awaiting_review",
        "archived": False,
    }])
    workbook = load_workbook(BytesIO(workbook_bytes))
    sheet = workbook["Corrective Actions"]

    assert sheet.freeze_panes == "A2"
    assert sheet["A2"].value is None
    assert sheet["B2"].value == "Glass audit"
    assert sheet["G2"].value == "Chris"
    assert sheet["H2"].value == "Mike"
    assert sheet["L2"].value == "Awaiting Review"
    assert "CorrectiveActionRegister" in sheet.tables


def test_action_workbook_protects_formula_like_text():
    workbook = load_workbook(BytesIO(build_action_workbook([{
        "audit_name": "=HYPERLINK(\"bad\")",
        "status": "open",
    }])))
    assert workbook["Corrective Actions"]["B2"].value.startswith("'=")
